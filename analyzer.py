
"""
ResumeIQ Pro - Resume Analysis Engine

Features:
- PDF and DOCX text extraction
- Skill recognition
- TF-IDF text similarity
- Skill coverage analysis
- Resume section detection
- Resume health checks
- Actionable recommendations
"""

import re
from io import BytesIO

import pymupdf
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# ------------------------------------------
# SKILL CATEGORIES
# ------------------------------------------

SKILL_CATEGORIES = {
    "Programming": [
        "python", "java", "javascript",
        "typescript", "c++", "c#",
        "html", "css", "sql", "r"
    ],

    "Data & Artificial Intelligence": [
        "machine learning",
        "artificial intelligence",
        "deep learning",
        "natural language processing",
        "nlp", "data analysis",
        "data science", "pandas",
        "numpy", "scikit-learn",
        "tensorflow", "pytorch",
        "power bi", "tableau",
        "data visualization"
    ],

    "Software Development": [
        "react", "angular", "flask",
        "django", "fastapi",
        "streamlit", "git", "github",
        "docker", "api",
        "rest api", "agile", "scrum"
    ],

    "Databases & Cloud": [
        "mysql", "postgresql",
        "mongodb", "aws",
        "azure", "cloud computing",
        "database design"
    ],

    "Cybersecurity & Networking": [
        "cybersecurity", "networking",
        "network security",
        "penetration testing",
        "firewall", "linux",
        "risk assessment"
    ],

    "Design & Quality Assurance": [
        "ui/ux", "figma",
        "wireframing", "prototyping",
        "manual testing",
        "automation testing",
        "regression testing",
        "user acceptance testing",
        "software testing"
    ],

    "Business & Professional": [
        "business analysis",
        "project management",
        "communication",
        "problem solving",
        "leadership", "teamwork",
        "excel", "requirements gathering"
    ]
}


# ------------------------------------------
# SKILL ALIASES
# ------------------------------------------

SKILL_ALIASES = {
    "natural language processing": ["nlp"],
    "artificial intelligence": ["ai"],
    "machine learning": ["ml"],
    "user acceptance testing": ["uat"],
    "power bi": ["powerbi"],
    "javascript": ["java script"],
    "postgresql": ["postgres"],
    "scikit-learn": ["sklearn"],
    "ui/ux": ["ui ux", "user interface design"],
    "rest api": ["restful api", "restful apis"]
}


# ------------------------------------------
# DOCUMENT EXTRACTION
# ------------------------------------------

def extract_resume_text(uploaded_file):
    """
    Extract readable text from PDF or DOCX.
    Supports Streamlit UploadedFile objects.
    """

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if len(file_bytes) > 10 * 1024 * 1024:
        raise ValueError(
            "File is too large. Maximum size is 10 MB."
        )

    if file_name.endswith(".pdf"):

        try:
            with pymupdf.open(
                stream=file_bytes,
                filetype="pdf"
            ) as document:

                if document.page_count > 30:
                    raise ValueError(
                        "PDF exceeds the 30-page limit."
                    )

                text = "\n".join(
                    page.get_text()
                    for page in document
                )

                return text

        except ValueError:
            raise
        except Exception as error:
            raise ValueError(
                "Unable to read the PDF. "
                "Check that it is valid and not encrypted."
            ) from error

    if file_name.endswith(".docx"):

        try:
            document = Document(BytesIO(file_bytes))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)

            return "\n".join(paragraphs)

        except Exception as error:
            raise ValueError(
                "Unable to read the Word document."
            ) from error

    raise ValueError(
        "Unsupported file type. Use PDF or DOCX."
    )


# ------------------------------------------
# SKILL RECOGNITION
# ------------------------------------------

def skill_is_present(text, skill):
    """Match a skill using word boundaries and aliases."""

    variations = [skill]
    variations.extend(SKILL_ALIASES.get(skill, []))

    for variation in variations:

        pattern = (
            r"(?<!\w)"
            + re.escape(variation)
            + r"(?!\w)"
        )

        if re.search(pattern, text, flags=re.IGNORECASE):
            return True

    return False


def extract_skills(text):
    """Return unique recognized skills."""

    found = set()

    for skills in SKILL_CATEGORIES.values():
        for skill in skills:
            if skill_is_present(text, skill):
                found.add(skill)

    return sorted(found)


def get_skill_categories(skills):
    """Group detected skills by category."""

    grouped = {}

    for category, category_skills in SKILL_CATEGORIES.items():

        matches = [
            skill
            for skill in skills
            if skill in category_skills
        ]

        if matches:
            grouped[category] = matches

    return grouped


# ------------------------------------------
# NLP TEXT SIMILARITY
# ------------------------------------------

def calculate_similarity(resume_text, job_text):
    """Calculate TF-IDF cosine similarity."""

    if not resume_text.strip() or not job_text.strip():
        return 0.0

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2)
        )

        vectors = vectorizer.fit_transform(
            [resume_text, job_text]
        )

        score = cosine_similarity(
            vectors[0:1],
            vectors[1:2]
        )[0][0]

        return round(float(score) * 100, 2)

    except ValueError:
        return 0.0


# ------------------------------------------
# RESUME SECTION CHECKS
# ------------------------------------------

SECTION_PATTERNS = {
    "Contact Information": [
        r"\b[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}\b"
    ],

    "Professional Summary": [
        r"\bprofessional summary\b",
        r"\bcareer objective\b",
        r"\bprofile summary\b",
        r"\babout me\b"
    ],

    "Work Experience": [
        r"\bwork experience\b",
        r"\bprofessional experience\b",
        r"\bemployment history\b",
        r"\binternship\b"
    ],

    "Education": [
        r"\beducation\b",
        r"\bacademic qualifications\b",
        r"\bqualifications\b"
    ],

    "Skills": [
        r"\btechnical skills\b",
        r"\bcore competencies\b",
        r"\bskills\b"
    ],

    "Projects": [
        r"\bprojects\b",
        r"\bacademic projects\b",
        r"\bpersonal projects\b"
    ]
}


def check_resume_sections(resume_text):
    """Check for common resume sections or signals."""

    results = {}

    for section, patterns in SECTION_PATTERNS.items():

        found = any(
            re.search(
                pattern,
                resume_text,
                flags=re.IGNORECASE
            )
            for pattern in patterns
        )

        results[section] = found

    return results


# ------------------------------------------
# RESUME HEALTH CHECK
# ------------------------------------------

def calculate_resume_health(resume_text):
    """
    Calculate a basic completeness score.

    This is a heuristic score, not an official
    ATS score or recruiter assessment.
    """

    sections = check_resume_sections(resume_text)

    found_count = sum(sections.values())
    total_sections = len(sections)

    score = round(
        (found_count / total_sections) * 100,
        2
    )

    return {
        "score": score,
        "sections": sections,
        "detected_count": found_count,
        "total_sections": total_sections
    }


# ------------------------------------------
# RECOMMENDATIONS
# ------------------------------------------

def generate_recommendations(
    missing_skills,
    resume_health
):
    """Generate actionable resume suggestions."""

    recommendations = []

    for skill in missing_skills[:5]:
        recommendations.append(
            f"If you have experience with {skill.title()}, "
            f"consider showing it clearly in your resume. "
            f"Otherwise, consider learning the skill."
        )

    for section, detected in (
        resume_health["sections"].items()
    ):
        if not detected:
            recommendations.append(
                f"Review your {section} section. "
                f"It was not recognized by the "
                f"basic resume structure checker."
            )

    if not recommendations:
        recommendations.append(
            "Your resume includes the recognized "
            "sections and job-related skills. "
            "Review it for clarity and accuracy."
        )

    return recommendations


# ------------------------------------------
# COMPLETE RESUME ANALYSIS
# ------------------------------------------

def analyze_resume(resume_text, job_text):
    """Return the full ResumeIQ Pro analysis."""

    resume_skills = extract_skills(resume_text)
    job_skills = extract_skills(job_text)

    matched_skills = sorted(
        set(resume_skills) & set(job_skills)
    )

    missing_skills = sorted(
        set(job_skills) - set(resume_skills)
    )

    similarity = calculate_similarity(
        resume_text,
        job_text
    )

    if job_skills:
        skill_coverage = round(
            len(matched_skills)
            / len(job_skills) * 100,
            2
        )
    else:
        skill_coverage = 0.0

    resume_health = calculate_resume_health(
        resume_text
    )

    recommendations = generate_recommendations(
        missing_skills,
        resume_health
    )

    return {
        "similarity": similarity,
        "skill_coverage": skill_coverage,
        "resume_skills": resume_skills,
        "job_skills": job_skills,
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "resume_health": resume_health,
        "skill_categories": get_skill_categories(
            resume_skills
        ),
        "recommendations": recommendations
    }
